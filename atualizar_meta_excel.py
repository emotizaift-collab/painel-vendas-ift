"""Busca Insights da Meta e cria dashboard_meta.xlsx.

Instalação:
    python -m pip install pandas openpyxl requests

Execução:
    python atualizar_meta_excel.py

As credenciais também podem ser definidas pelas variáveis de ambiente
META_ACCESS_TOKEN e META_AD_ACCOUNT_ID.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ===================== PREENCHA SUAS CREDENCIAIS =====================
ACCESS_TOKEN = os.getenv("META_ACCESS_TOKEN", "COLE_SEU_ACCESS_TOKEN_AQUI")
AD_ACCOUNT_ID = os.getenv("META_AD_ACCOUNT_ID", "act_COLE_SEU_AD_ACCOUNT_ID_AQUI")
# ======================================================================

API_VERSION = "v20.0"
OUTPUT_FILE = Path(__file__).with_name("dashboard_meta.xlsx")
BLUE = "1E3A8A"
LIGHT_ROW = "F8FAFC"
WHITE = "FFFFFF"
GRAY = "64748B"
THIN_GRAY = Side(style="thin", color="E2E8F0")


def validate_credentials() -> None:
    if not ACCESS_TOKEN or ACCESS_TOKEN == "COLE_SEU_ACCESS_TOKEN_AQUI":
        raise ValueError("Preencha ACCESS_TOKEN com um token válido da Meta.")
    if not AD_ACCOUNT_ID or "COLE_SEU_AD_ACCOUNT_ID" in AD_ACCOUNT_ID:
        raise ValueError("Preencha AD_ACCOUNT_ID com o ID da conta, incluindo act_.")


def action_value(actions: list[dict[str, Any]] | None, action_type: str) -> float:
    """Soma o valor da ação solicitada; ausência da ação equivale a zero."""
    return sum(
        float(action.get("value", 0) or 0)
        for action in (actions or [])
        if action.get("action_type") == action_type
    )


def fetch_insights() -> pd.DataFrame:
    validate_credentials()
    url = f"https://graph.facebook.com/{API_VERSION}/{AD_ACCOUNT_ID}/insights"
    params = {
        "access_token": ACCESS_TOKEN,
        "level": "campaign",
        "time_increment": 1,
        # A Graph API usa "last_30d" (e não "last_30days") para este preset.
        "date_preset": "last_30d",
        "fields": "date_start,campaign_id,campaign_name,spend,impressions,clicks,actions",
        "limit": 100,
    }
    records: list[dict[str, Any]] = []
    next_url: str | None = url
    next_params: dict[str, Any] | None = params
    while next_url:
        response = requests.get(next_url, params=next_params, timeout=60)
        if not response.ok:
            try:
                error = response.json().get("error", {})
            except ValueError:
                error = response.text
            raise RuntimeError(f"Erro da API Meta ({response.status_code}): {error}")
        payload = response.json()
        records.extend(payload.get("data", []))
        next_url = payload.get("paging", {}).get("next")
        next_params = None

    rows = [
        {
            "Data": record.get("date_start"),
            "ID Campanha": record.get("campaign_id", ""),
            "Campanha": record.get("campaign_name", "Sem nome"),
            "Gasto": float(record.get("spend", 0) or 0),
            "Impressões": int(float(record.get("impressions", 0) or 0)),
            "Cliques": int(float(record.get("clicks", 0) or 0)),
            "Leads": action_value(record.get("actions"), "lead"),
            "Vendas": action_value(record.get("actions"), "purchase"),
        }
        for record in records
    ]
    frame = pd.DataFrame(
        rows,
        columns=[
            "Data", "ID Campanha", "Campanha", "Gasto",
            "Impressões", "Cliques", "Leads", "Vendas",
        ],
    )
    if frame.empty:
        raise RuntimeError("A API não retornou dados para os últimos 30 dias.")
    frame["Data"] = pd.to_datetime(frame["Data"])
    return frame.sort_values("Data").reset_index(drop=True)


def add_table(
    sheet: Any, name: str, end_row: int, end_col: int, start_row: int = 1
) -> None:
    table = Table(
        ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}",
        displayName=name,
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def campaign_sheet_title(campaign_name: str, used_titles: set[str]) -> str:
    """Cria um nome de aba válido e único a partir do nome da campanha."""
    invalid = set(r'[]:*?/\\')
    base = "".join("_" if character in invalid else character for character in campaign_name)
    base = base.strip() or "Sem nome"
    base = base[:31]
    title = base
    suffix = 2
    while title in used_titles:
        suffix_text = f"_{suffix}"
        title = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_titles.add(title)
    return title


def create_campaign_sheet(workbook: Workbook, campaign_name: str, campaign_data: pd.DataFrame) -> None:
    """Cria uma visão individual de campanha/produto."""
    sheet = workbook.create_sheet(campaign_sheet_title(campaign_name, {item.title for item in workbook.worksheets}))
    sheet.sheet_view.showGridLines = True
    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"Produto/Campanha: {campaign_name}"
    sheet["A1"].font = Font(name="Segoe UI", bold=True, size=20, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet["A2"] = "Últimos 30 dias | Dados da Meta Ads"
    sheet["A2"].font = Font(name="Segoe UI", italic=True, color=GRAY)

    metrics = [
        ("A3:B3", "A4:B4", "Gasto", campaign_data["Gasto"].sum(), 'R$ #,##0.00'),
        ("C3:D3", "C4:D4", "Leads", campaign_data["Leads"].sum(), '#,##0.00'),
        ("E3:F3", "E4:F4", "Vendas", campaign_data["Vendas"].sum(), '#,##0.00'),
        ("G3:H3", "G4:H4", "CAC", '=IFERROR(A4/E4,0)', 'R$ #,##0.00'),
    ]
    for label_range, value_range, label, value, number_format in metrics:
        sheet.merge_cells(label_range)
        sheet.merge_cells(value_range)
        label_cell = sheet[label_range.split(":")[0]]
        value_cell = sheet[value_range.split(":")[0]]
        label_cell.value = label
        label_cell.font = Font(name="Segoe UI", bold=True, color=WHITE)
        label_cell.fill = PatternFill("solid", fgColor=BLUE)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.value = value
        value_cell.font = Font(name="Segoe UI", bold=True, size=16, color=BLUE)
        value_cell.fill = PatternFill("solid", fgColor=LIGHT_ROW)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="center")
        for cell_range in (label_range, value_range):
            for row in sheet[cell_range]:
                for cell in row:
                    cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)

    sheet["A7"] = "Desempenho diário"
    sheet["A7"].font = Font(name="Segoe UI", bold=True, size=14, color=BLUE)
    headers = ["Data", "Gasto", "Leads", "Vendas"]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(8, column, header)
        cell.font = Font(name="Segoe UI", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
    daily = (
        campaign_data.groupby("Data", as_index=False)[["Gasto", "Leads", "Vendas"]]
        .sum()
        .sort_values("Data")
    )
    for row_number, row in enumerate(daily.itertuples(index=False, name=None), 9):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value)
        sheet.cell(row_number, 1).number_format = "dd/mm/yyyy"
        sheet.cell(row_number, 2).number_format = 'R$ #,##0.00'
        sheet.cell(row_number, 3).number_format = '#,##0.00'
        sheet.cell(row_number, 4).number_format = '#,##0.00'
        if row_number % 2 == 1:
            for cell in sheet[row_number][:4]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_ROW)
    if not daily.empty:
        add_table(sheet, f"TabelaCampanha{sheet.max_row}", sheet.max_row, 4, start_row=8)
        chart = LineChart()
        chart.title = "Gasto vs Vendas"
        chart.style = 13
        chart.y_axis.title = "Valor"
        chart.x_axis.title = "Data"
        chart.height = 8
        chart.width = 14
        chart.add_data(
            Reference(sheet, min_col=2, max_col=2, min_row=8, max_row=sheet.max_row),
            titles_from_data=True,
        )
        chart.add_data(
            Reference(sheet, min_col=4, max_col=4, min_row=8, max_row=sheet.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(sheet, min_col=1, min_row=9, max_row=sheet.max_row))
        sheet.add_chart(chart, "F7")
    for column, width in enumerate([14, 16, 14, 14, 16, 16, 16, 16], 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A9"


def style_sheet(sheet: Any) -> None:
    sheet.sheet_view.showGridLines = True
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Segoe UI", color="0F172A")
            cell.border = Border(bottom=THIN_GRAY)
    for cell in sheet[1]:
        cell.font = Font(name="Segoe UI", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center")
    for row_number in range(2, sheet.max_row + 1, 2):
        for cell in sheet[row_number]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_ROW)


def create_workbook(data: pd.DataFrame) -> None:
    workbook = Workbook()
    raw = workbook.active
    raw.title = "Dados_Brutos"
    raw.append(list(data.columns))
    for row in data.itertuples(index=False, name=None):
        raw.append(list(row))
    for row in raw.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[3].number_format = 'R$ #,##0.00'
        for cell in row[4:]:
            cell.number_format = "#,##0.00"
    style_sheet(raw)
    raw.freeze_panes = "A2"
    widths = [14, 18, 28, 15, 16, 13, 12, 12]
    for index, width in enumerate(widths, 1):
        raw.column_dimensions[get_column_letter(index)].width = width
    add_table(raw, "TabelaDadosBrutos", raw.max_row, raw.max_column)

    dashboard = workbook.create_sheet("Dashboard")
    dashboard.sheet_view.showGridLines = True
    dashboard.merge_cells("A1:J1")
    dashboard["A1"] = "Dashboard Meta Ads"
    dashboard["A1"].font = Font(name="Segoe UI", bold=True, size=22, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=BLUE)
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")
    dashboard.row_dimensions[1].height = 34
    dashboard["A2"] = "Últimos 30 dias | Atualizado diretamente pela Meta Insights API"
    dashboard["A2"].font = Font(name="Segoe UI", italic=True, color=GRAY)

    raw_end = raw.max_row
    campaign_summary = (
        data.groupby(["ID Campanha", "Campanha"], as_index=False)
        .agg(
            Gasto=("Gasto", "sum"),
            Impressões=("Impressões", "sum"),
            Cliques=("Cliques", "sum"),
            Leads=("Leads", "sum"),
            Vendas=("Vendas", "sum"),
        )
        .sort_values("Gasto", ascending=False)
    )
    cards = [
        ("A3:B3", "A4:B4", "Total Gasto", f"=SUM(Dados_Brutos!D:D)", 'R$ #,##0.00'),
        ("C3:D3", "C4:D4", "Total Leads", f"=SUM(Dados_Brutos!G:G)", '#,##0.00'),
        ("E3:F3", "E4:F4", "Total Vendas", f"=SUM(Dados_Brutos!H:H)", '#,##0.00'),
        ("G3:H3", "G4:H4", "Custo por Lead (CPL)", '=IFERROR(A4/C4,0)', 'R$ #,##0.00'),
        ("I3:J3", "I4:J4", "Custo por Venda (CAC)", '=IFERROR(A4/E4,0)', 'R$ #,##0.00'),
    ]
    for label_range, value_range, label, formula, number_format in cards:
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(":")[0]]
        value_cell = dashboard[value_range.split(":")[0]]
        label_cell.value = label
        label_cell.font = Font(name="Segoe UI", bold=True, color=WHITE)
        label_cell.fill = PatternFill("solid", fgColor=BLUE)
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.value = formula
        value_cell.font = Font(name="Segoe UI", bold=True, size=16, color=BLUE)
        value_cell.fill = PatternFill("solid", fgColor=LIGHT_ROW)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="center")
        for cell_range in (label_range, value_range):
            for row in dashboard[cell_range]:
                for cell in row:
                    cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
    dashboard.row_dimensions[3].height = 24
    dashboard.row_dimensions[4].height = 34

    dashboard["A7"] = "Resumo por campanha"
    dashboard["A7"].font = Font(name="Segoe UI", bold=True, size=14, color=BLUE)
    campaign_headers = ["ID Campanha", "Campanha", "Gasto", "Leads", "Vendas", "CPL", "CAC"]
    for column, header in enumerate(campaign_headers, 1):
        dashboard.cell(8, column, header)
    for index, row in enumerate(campaign_summary.itertuples(index=False, name=None), 9):
        dashboard.cell(index, 1, row[0])
        dashboard.cell(index, 2, row[1])
        dashboard.cell(index, 3, f'=SUMIF(Dados_Brutos!C:C,B{index},Dados_Brutos!D:D)').number_format = 'R$ #,##0.00'
        dashboard.cell(index, 4, f'=SUMIF(Dados_Brutos!C:C,B{index},Dados_Brutos!G:G)').number_format = '#,##0.00'
        dashboard.cell(index, 5, f'=SUMIF(Dados_Brutos!C:C,B{index},Dados_Brutos!H:H)').number_format = '#,##0.00'
        dashboard.cell(index, 6, f'=IFERROR(C{index}/D{index},0)').number_format = 'R$ #,##0.00'
        dashboard.cell(index, 7, f'=IFERROR(C{index}/E{index},0)').number_format = 'R$ #,##0.00'
    for cell in dashboard[8][:3]:
        cell.font = Font(name="Segoe UI", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
    summary_end = 8 + len(campaign_summary)
    for row_number in range(9, summary_end + 1, 2):
        for cell in dashboard[row_number][:7]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_ROW)
    add_table(dashboard, "TabelaCampanhas", summary_end, 7, start_row=8)
    dashboard.freeze_panes = "A9"

    chart = LineChart()
    chart.title = "Gasto vs Vendas por campanha"
    chart.style = 13
    chart.y_axis.title = "Valor"
    chart.x_axis.title = "Campanha"
    chart.height = 9
    chart.width = 18
    chart.add_data(Reference(dashboard, min_col=3, max_col=3, min_row=8, max_row=summary_end), titles_from_data=True)
    chart.add_data(Reference(dashboard, min_col=5, max_col=5, min_row=8, max_row=summary_end), titles_from_data=True)
    chart.set_categories(Reference(dashboard, min_col=2, min_row=9, max_row=summary_end))
    dashboard.add_chart(chart, "I7")

    sales = workbook.create_sheet("Vendas")
    sales_data = data.loc[
        data["Vendas"] > 0,
        ["Data", "ID Campanha", "Campanha", "Vendas", "Gasto"],
    ].copy()
    sales_data = sales_data.rename(columns={"Vendas": "Quantidade de Vendas"})
    sales.append(list(sales_data.columns))
    for row in sales_data.itertuples(index=False, name=None):
        sales.append(list(row))
    for row in sales.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[3].number_format = "#,##0.00"
        row[4].number_format = 'R$ #,##0.00'
    style_sheet(sales)
    sales.freeze_panes = "A2"
    for index, width in enumerate([14, 18, 28, 20, 15], 1):
        sales.column_dimensions[get_column_letter(index)].width = width
    if sales.max_row > 1:
        add_table(sales, "TabelaVendas", sales.max_row, sales.max_column)
    for column in range(1, 17):
        dashboard.column_dimensions[get_column_letter(column)].width = 16
    dashboard.column_dimensions["B"].width = 28
    dashboard.column_dimensions["D"].width = 3
    for campaign_name, campaign_data in data.groupby("Campanha", sort=True):
        create_campaign_sheet(workbook, campaign_name, campaign_data)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(OUTPUT_FILE)


def main() -> None:
    data = fetch_insights()
    create_workbook(data)
    print(f"Arquivo criado com sucesso: {OUTPUT_FILE}")
    print(f"Registros processados: {len(data)}")


if __name__ == "__main__":
    main()
